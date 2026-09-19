import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)  # Slate dark

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("GridVision OpenDSS Studio — LECO BZ0109\nTechnical Manual & Electrical Calculation Specification")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Comprehensive Technical Equations, Input File Data Mapping, Engineering Assumptions, and QSTS Power Flow Methodology")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)  # Primary Accent Blue
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        return p

    def add_equation_box(eq_text, explanation_text):
        p_eq = doc.add_paragraph()
        p_eq.paragraph_format.left_indent = Inches(0.25)
        p_eq.paragraph_format.right_indent = Inches(0.25)
        p_eq.paragraph_format.space_before = Pt(6)
        p_eq.paragraph_format.space_after = Pt(2)
        run_eq = p_eq.add_run(eq_text)
        run_eq.font.name = 'Consolas'
        run_eq.font.size = Pt(11)
        run_eq.font.bold = True
        run_eq.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)

        p_exp = doc.add_paragraph()
        p_exp.paragraph_format.left_indent = Inches(0.25)
        p_exp.paragraph_format.right_indent = Inches(0.25)
        p_exp.paragraph_format.space_after = Pt(8)
        run_exp = p_exp.add_run(explanation_text)
        run_exp.font.name = 'Calibri'
        run_exp.font.size = Pt(10)
        run_exp.font.italic = True
        run_exp.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # -------------------------------------------------------------
    # 1. EXECUTIVE SUMMARY & OVERVIEW
    # -------------------------------------------------------------
    add_heading_1("1. Executive Summary & Overview")
    p = doc.add_paragraph(
        "This document provides the complete engineering specification for the GridVision OpenDSS Studio web application. "
        "The software performs Quasi-Static Time Series (QSTS) 3-phase 4-wire power flow simulations on the LECO Transformer BZ0109 low-voltage (400V/230V) distribution network. "
        "It integrates monthly billing consumption datasets, 15-minute smart meter Load Profiling (LP) interval records, rooftop solar PV inverter reports, and GIS geographic network topology."
    )
    p.paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # 2. RAW DATA INGESTION & FILE MAPPING
    # -------------------------------------------------------------
    add_heading_1("2. Raw Data Ingestion & Input File Lineage")
    p = doc.add_paragraph(
        "The software extracts customer demand, solar capacity, network topology, and smart meter measurements from specific input files, sheets, and columns as detailed below:"
    )
    p.paragraph_format.space_after = Pt(8)

    table_data = doc.add_table(rows=1, cols=4)
    table_data.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table_data.rows[0].cells
    headers = ["Data Category", "File Name & Path", "Sheet Name & Column Headers", "Extracted Parameters & Rules"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], '0F172A')
        for p in hdr_cells[i].paragraphs:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            p.runs[0].font.size = Pt(9.5)

    data_mapping_rows = [
        ("GIS & Topology", "network_map.json", "Root JSON object:\n- poles[]\n- lines[]", "Poles: id, lat, lon, is_tf (AR48T), feeder (F1–F8).\nLines: from, to, length_m, conductor (ABC_70, Fly_AAC)."),
        ("Billing & Demands", "DATA_BZ0109.xlsx /\naccount_pole_map.json", "Sheet: DATA / Consumption\nColumns: ACCOUNT, POLE ID, CATEGORY / TARIFF, kWh columns", "Account matching rule: strip leading zeros (e.g., '00024/1234' -> '24/1234'). Extracts monthly energy E_month and tariff class (R1, GP11, STL, I12)."),
        ("Smart Meter LP", "pole_voltage_compare.json /\nvoltages_timeline.json", "Sheet: LP Data\nColumns: CUSTOMER_REF, DATE, TIME, AVG._IMPORT_KW, AVG._EXPORT_KW, PHASE_A_VOLTAGE, PHASE_B_VOLTAGE, PHASE_C_VOLTAGE", "15-minute interval transformer net power P_meas(t) and pole measured phase voltages (Va, Vb, Vc). Used for calibration and measured vs. simulated comparison."),
        ("Rooftop Solar PV", "BZ0109_solar_customers.csv /\nSolar Data V2", "Sheet: Solar Data V2\nColumns: ACCOUNT / CUSTOMER, POLE ID, INVERTER_KW / CAPACITY", "Extracts solar PV inverter capacity P_inv (kW) mapped to customer account and pole ID (11 accounts, 179.3 kW total capacity)."),
    ]

    for cat, fname, sheet, rules in data_mapping_rows:
        row_cells = table_data.add_row().cells
        row_cells[0].text = cat
        row_cells[1].text = fname
        row_cells[2].text = sheet
        row_cells[3].text = rules
        for c in row_cells:
            set_cell_background(c, 'F8FAFC')
            for p in c.paragraphs:
                p.runs[0].font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 3. SCIENTIFIC & ENGINEERING ASSUMPTIONS
    # -------------------------------------------------------------
    add_heading_1("3. Scientific & Engineering Assumptions")
    assumptions = [
        ("Nominal System Voltage", "Nominal Phase-to-Neutral voltage V_nom = 230.0 V (Phase-to-Phase V_LL = 400.0 V) at 50 Hz system frequency."),
        ("Statutory Voltage Limits", "Statutory allowable voltage limits strictly enforced at 216.2 V (-6.0%) minimum under-voltage threshold and 243.8 V (+6.0%) maximum over-voltage threshold."),
        ("Transformer Rating", "Distribution Transformer Tx_BZ0109 rated at 250 kVA, 11kV/0.4kV, Delta-Wye grounded configuration (%R = 1.0%, %X = 4.5%)."),
        ("Customer Power Factor", "Standard customer loads assigned a lagging power factor cos φ = 0.95 (tan φ = 0.329). Solar PV inverters operate at unity power factor (PF = 1.0)."),
        ("Radial Tree Topology", "The distribution network is modeled as a strict radial tree graph originating from transformer pole AR48T across 8 main distributor branches (F1 through F8)."),
        ("Quasi-Static Time Series", "QSTS evaluates 96 steady-state snapshot power flow solves per day (15-minute time steps) assuming fast transient dynamics settle within each interval."),
        ("Behind-the-Meter Netting", "Smart meters record NET power flow (Import - Export). Internal home customer load behind the solar meter consumes a portion of solar output before reaching the meter."),
        ("Conductor Impedances", "Aerial Bundled Conductors (ABC_70: R1 = 0.443 Ω/km, X1 = 0.09 Ω/km) and All Aluminum Conductors (Fly_AAC: R1 = 0.270 Ω/km, X1 = 0.30 Ω/km)."),
    ]

    table_assump = doc.add_table(rows=1, cols=2)
    table_assump.alignment = WD_TABLE_ALIGNMENT.CENTER
    hcells = table_assump.rows[0].cells
    hcells[0].text = "Engineering Parameter"
    hcells[1].text = "Scientific Assumption & Justification"
    for i, h in enumerate([hcells[0], hcells[1]]):
        set_cell_background(h, '0F172A')
        for p in h.paragraphs:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            p.runs[0].font.size = Pt(9.5)

    for param, desc in assumptions:
        rcells = table_assump.add_row().cells
        rcells[0].text = param
        rcells[1].text = desc
        for c in rcells:
            set_cell_background(c, 'F8FAFC')
            for p in c.paragraphs:
                p.runs[0].font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 4. TECHNICAL & MATHEMATICAL EQUATIONS
    # -------------------------------------------------------------
    add_heading_1("4. Technical & Mathematical Equations")

    add_heading_2("4.1. Customer Demand Allocation")
    add_equation_box(
        "P_avg,c = E_month,c / (N_days × 24)",
        "Where E_month,c is the monthly billed kWh for account c from DATA_BZ0109.xlsx, N_days is the billing cycle days (e.g., 30 days). Converts energy into mean active power (kW)."
    )

    add_heading_2("4.2. Tariff-Shaped Customer Active & Reactive Load")
    add_equation_box(
        "P_c(t) = P_avg,c × s_tariff(c)(t) × α(t)\nQ_c(t) = P_c(t) × tan(acos(PF))",
        "Where s_tariff(c)(t) is the normalized 15-minute daily tariff profile vector for tariff class c (R1, GP11, STL, I12), α(t) is the hourly calibration multiplier, and PF = 0.95."
    )

    add_heading_2("4.3. 15-Minute Solar Irradiance Curve & PV Output")
    add_equation_box(
        "c_irr(t) = max(0.0, 1.0 - ((h(t) - 12.25) / 5.5)^2)  for 6.0 <= h(t) <= 18.0\nP_solar,k(t) = P_inv,k × c_irr(t) × β",
        "Where h(t) = t/4.0 is hour of day (0.0 to 23.75), P_inv,k is installed inverter kW rating at pole k, c_irr(t) is normalized irradiance (0.0 at night, 1.0 at 12:15 solar noon), and β is irradiance scale."
    )

    add_heading_2("4.4. Locked Export-Target Solar Calibration")
    add_equation_box(
        "P_pv*(t) = max(P_load(t) - P_meas(t), 0.0)",
        "Where P_meas(t) is the measured transformer net power from LP data. Used in locked calibration mode so simulated net power tracks physical meter export when solar output exceeds load."
    )

    add_heading_2("4.5. Global Active Power Conservation")
    add_equation_box(
        "P_TF,net(t) = P_load(t) - P_pv(t) + P_loss(t)",
        "Where P_TF,net > 0 represents net power import from the 11kV grid through the 250kVA transformer, and P_TF,net < 0 represents net solar power export back to the 11kV grid."
    )

    add_heading_2("4.6. OpenDSS 3-Phase 4-Wire Branch Voltage Drop")
    add_equation_box(
        "[V_a, V_b, V_c, V_n]_(k+1) = [V_a, V_b, V_c, V_n]_k - Z_4x4 × [I_a, I_b, I_c, I_n]_k",
        "Where Z_4x4 is the 4-wire line impedance matrix derived from conductor resistance R and self/mutual reactance X for ABC/Fly overhead line codes."
    )

    add_heading_2("4.7. Phase Voltage Averages & Unbalance Percentage")
    add_equation_box(
        "V_avg = (V_a + V_b + V_c) / 3.0\nUnbalance % = ((max(V_a,V_b,V_c) - min(V_a,V_b,V_c)) / V_avg) × 100%",
        "Calculates mean phase-to-neutral voltage V_avg and phase voltage unbalance percentage at any pole node across 3-phase conductors (lines 3485-3488 in index.html)."
    )

    add_heading_2("4.8. Transformer Loading Percentage")
    add_equation_box(
        "Loading % = (|P_TF,net(t)| / S_rated) × 100%",
        "Where S_rated = 250 kVA. Measures transformer capacity utilization under import peak (evening) and export peak (solar noon)."
    )

    add_heading_2("4.9. Statistical Model Accuracy Metrics (RMSE & Pearson Correlation)")
    add_equation_box(
        "RMSE_kW = sqrt((1/N) × Σ (P_sim(t) - P_meas(t))^2)\nRMSE_V = sqrt((1/M) × Σ (V_sim(k) - V_meas(k))^2)\nρ = Σ (P_meas - P_bar_meas)(P_sim - P_bar_sim) / [ sqrt(Σ (P_meas - P_bar_meas)^2) × sqrt(Σ (P_sim - P_bar_sim)^2) ]",
        "Quantifies simulation accuracy against physical LP smart meter measurements for active power (kW), voltage (V), and daily shape correlation (ρ)."
    )

    # -------------------------------------------------------------
    # 5. EXECUTION & SIMULATION WORKFLOW
    # -------------------------------------------------------------
    add_heading_1("5. Execution & Simulation Workflow")
    p_wf = doc.add_paragraph(
        "1. Data Reading: FastAPI backend reads GIS coordinates, billing customer accounts, LP 15-minute readings, and solar inverter reports.\n"
        "2. Topology Generation: BFS algorithm constructs the radial feeder graph from transformer pole AR48T and calculates cumulative electrical distance d(k) along line paths.\n"
        "3. OpenDSS Script Generator: dss_generator.py writes master.dss, lines.dss, loads.dss, and pvsystem.dss files.\n"
        "4. QSTS Solve Loop: opendss_runner.py steps sequentially through 96 15-minute intervals, solving nodal voltages, current flows, transformer loading, and line losses.\n"
        "5. Result Visualization: The web browser displays 24-hour time-series charts, Leaflet voltage heat maps, measured vs. simulated phase cards, and distance vs. voltage feeder profiles."
    )
    p_wf.paragraph_format.space_after = Pt(12)

    # Save output
    out_path = Path("c:/Users/Gavesh/Downloads/OPENDSS/GridVision_OpenDSS_Technical_Report.docx")
    doc.save(str(out_path))
    print(f"Report saved successfully to {out_path}")

if __name__ == "__main__":
    create_report()
